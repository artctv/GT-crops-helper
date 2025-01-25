import json
from pathlib import Path
from dataclasses import dataclass
from dataclasses import astuple

import pandas as pd
from openpyxl import load_workbook


basedir = Path(__file__).parent.resolve()
print(basedir)
sheet_path = basedir / "gt_crops_stats.xlsx"
json_path = basedir / "crops.json"
data_path = basedir / "data"
data_path.mkdir(exist_ok=True)
min_row = 2
max_row = 72
max_col = 21

@dataclass(eq=True, frozen=True)
class CropProbStats:
    che: int
    foo: int
    def_: int
    col: int
    wee: int


@dataclass(eq=True, frozen=True)
class CropStats:
    id_: int
    name: str
    tier: int
    attributes: str
    mod: str
    stats: CropProbStats

    def __repr__(self):
        return self.name
    
    def __int__(self):
        return self.id_


def parse_data():
    data: list[CropStats] = []
    with open(json_path, "r") as f:
        # d = json.load(f)
        for i in json.load(f):
            v = i["stats"].pop("def")
            i["stats"]["def_"] = v
            prob_stat = CropProbStats(
                **i["stats"]
            )
            i.pop("stats")
            i.pop("rus")
            v = i.pop("id")
            i["id_"] = v
            crop = CropStats(
                **i,
                stats=prob_stat
            )
            data.append(crop)

    return data


crops_json = parse_data()

heads = [
    "InternalId",
    "Code",
    "Name",
    "Tier",
    "Size",
    "Speed",
    "aHarvest",
    "Harvest",
    "Chemistry",
    "Food",
    "Defensive",
    "Colorful",
    "Weed",
    "RequiredBlock",
    "RootsLength",
    "GrowDuaration",
    "MaxSumOfGrow",
    "MinSumOfGrow",
    "GainChance",
    "OptimalHarvestSize",
    "AdditionalRequirements",
    "WeightInfluences"
]
attrs_heads = [
    "InternalId",
    "Attribute"
]
mtm_attrs_heads = [
    "CropId",
    "AttributeId"
]

@dataclass
class Crop:
    Code: str
    Name: str
    Tier: int
    Size: int
    Speed: int
    aHarvest: int | None
    Harvest: int | None
    Chemistry: int
    Food: int
    Defensive: int
    Colorful: int
    Weed: int
    RequiredBlock: str
    RootsLength: int
    GrowDuaration: int | str
    MaxSumOfGrow: int
    MinSumOfGrow: int | None
    GainChance: float
    OptimalHarvestSize: int
    AdditionalRequirements: str
    WeightInfluences: str

    def __post_init__(self):
        self.Name = self.Name[0].upper() + self.Name[1:]
        # self.Name = self.Name.title()



def get_sheet():
    wb = load_workbook(filename = sheet_path)
    sheet = wb["Лист1"]
    return sheet


def iter_rows(sheet):
    for col in sheet.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
        yield col


def build_head(sheet):
    for col in sheet.iter_cols(min_row=1, max_col=max_col, max_row=1):
        for cell in col:
            print(cell.value)


def get_json_crop(crop_name: str):
    for i in crops_json:
        if i.name == crop_name:
            return i


def parse_row(row):
    rows_data = [row[i].value for i in range(max_col)]
    crop = Crop(*rows_data)
    return astuple(crop)


def build_crop_attributes():
    attributes = []
    for i in crops_json:
        for k in i.attributes.split(" "):
            if k not in attributes:
                attributes.append(k)
    
    attributes = set(attributes)
    return {i:n for n, i in enumerate(attributes, start=1)}



def handle_data(sheet):
    crop_data, crop_attrs, crop_data_attrs_mtm = [], build_crop_attributes(), []
    for id_, i in enumerate(iter_rows(sheet), start=1):
        crop = parse_row(i)
        crop = [id_, *crop]
        crop_data.append(crop)
        json_crop = get_json_crop(crop_name=crop[2])
        crop_attr = json_crop.attributes
        for k in crop_attr.split(" "):
            attr_id = crop_attrs[k]
            crop_data_attrs_mtm.append(
                [id_, attr_id]
            )
    crop_df = pd.DataFrame(crop_data, columns=heads, index=None)
    crop_df.reset_index(drop=True, inplace=True)
    # crop_df["GainChance"] = crop_df["GainChance"].apply(lambda x: round(x, 3))
    crop_df.to_csv(path_or_buf=data_path / "crops.csv", sep=";", index=False)

    crop_attrs = [[v, k] for k, v in crop_attrs.items()]
    attrs_df = pd.DataFrame(crop_attrs, columns=attrs_heads, index=None,)
    attrs_df.reset_index(drop=True, inplace=True)
    attrs_df.to_csv(path_or_buf=data_path / "attributes.csv", sep=";", index=False)
    
    crop_attrs_mtm_df = pd.DataFrame(crop_data_attrs_mtm, columns=mtm_attrs_heads, index=None)
    crop_attrs_mtm_df.reset_index(drop=True, inplace=True)
    crop_attrs_mtm_df.to_csv(path_or_buf=data_path / "crops_attributes.csv", sep=";", index=False)


def main():
    sheet = get_sheet()
    handle_data(sheet)


if __name__ == "__main__":
    main()
